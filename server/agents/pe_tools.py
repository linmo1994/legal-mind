"""Plan-and-Execute whitelist tools (one call per plan step)."""

from __future__ import annotations

from typing import Any, Callable, Dict, List, Optional

TOOL_NAMES = (
    "retrieve_law",
    "retrieve_case",
    "read_evidence",
    "draft_doc",
    "reason",
)

OBS_MAX = 6000

_DOC_SKILL_AGENTS = frozenset({"doc_writing", "orchestrator"})

_DRAFT_DOC_SYSTEM = (
    "你是法律文书助手。根据用户指示起草完整法律文书正文。"
    "不要编造未提供的当事人信息；缺失处写「【待补充】」。"
    "只输出文书正文，不要输出 JSON、不要解释写作过程。"
    "结尾必须为文书式落款（例如起诉状：「此致」+ 人民法院 + 具状人 + 日期占位；"
    "协议类：签署栏/各方签章与日期占位），不得以聊天口吻收尾。"
)


def _skills_for_draft_doc(skills: Optional[List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for sk in skills or []:
        if not isinstance(sk, dict):
            continue
        applies = sk.get("applies_to") or []
        if not isinstance(applies, list):
            applies = [applies]
        if _DOC_SKILL_AGENTS.intersection({str(a) for a in applies}):
            out.append(sk)
    return out


def _draft_doc_system_prompt(skills: Optional[List[Dict[str, Any]]]) -> str:
    system = _DRAFT_DOC_SYSTEM
    blocks: List[str] = []
    for sk in _skills_for_draft_doc(skills):
        name = str(sk.get("name") or sk.get("id") or "技能").strip()
        body = str(sk.get("body") or "").strip()
        if not body:
            continue
        blocks.append(f"技能《{name}》\n{body}")
    if blocks:
        system += "\n【内部技能，禁止写入文书正文】\n" + "\n\n".join(blocks)
        system += "\n请严格按上述技能要求组织文书结构与落款。"
    return system


def _infer_draft_title(text: str) -> str:
    t = text or ""
    if "起诉状" in t:
        return "民事起诉状"
    if "答辩" in t:
        return "答辩状"
    if "协议" in t:
        return "协议书"
    if "判决" in t:
        return "模拟判决书"
    if "申请书" in t:
        return "申请书"
    return "法律文书"


def _load_template_docx_bytes(file_service: Any, file_id: Optional[str]) -> Optional[bytes]:
    if not file_service or not file_id:
        return None
    info = file_service.get_file(file_id)
    if not isinstance(info, dict):
        return None
    path = info.get("file_path")
    if path:
        try:
            with open(path, "rb") as f:
                return f.read()
        except Exception:
            pass
    get_bytes = getattr(file_service, "get_file_bytes", None)
    if callable(get_bytes):
        try:
            data = get_bytes(file_id)
            if data:
                return data
        except Exception:
            pass
    return None


def _artifact_from_saved_docx(
    info: Dict[str, Any],
    *,
    title: str,
    preview: str = "",
) -> Optional[Dict[str, Any]]:
    file_id = info.get("file_id")
    if not file_id:
        return None
    preview = (preview or "").strip()
    if len(preview) > 600:
        preview = preview[:600] + "…"
    filename = info.get("original_name") or f"{title}.docx"
    return {
        "filename": filename,
        "file_id": file_id,
        "download_url": f"/api/files/{file_id}/download",
        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "title": title,
        "preview": preview,
    }


def _preview_text_from_docx_bytes(docx_bytes: bytes, *, limit: int = 600) -> str:
    """Plain-text preview from a filled/exported docx (paragraphs + table cells)."""
    if not docx_bytes:
        return ""
    try:
        import io

        from docx import Document

        doc = Document(io.BytesIO(docx_bytes))
        parts: List[str] = []
        for para in doc.paragraphs:
            t = (para.text or "").strip()
            if t:
                parts.append(t)
        for table in doc.tables:
            for row in table.rows:
                seen = set()
                cells_out: List[str] = []
                for cell in row.cells:
                    cid = id(cell._tc)
                    if cid in seen:
                        continue
                    seen.add(cid)
                    t = (cell.text or "").strip()
                    if t:
                        cells_out.append(t.replace("\n", " "))
                if cells_out:
                    parts.append(" | ".join(cells_out))
        text = "\n".join(parts).strip()
        if len(text) > limit:
            return text[: limit - 1] + "…"
        return text
    except Exception as exc:
        print(f"[pe_tools] docx preview extract failed: {exc}")
        return ""


def _export_docx_artifact(
    title: str,
    body: str,
    file_service: Any,
    session_id: Optional[str] = None,
) -> Optional[Dict[str, Any]]:
    """Same shape as orchestrator doc_writing artifact; None if export unavailable."""
    if not file_service or not (body or "").strip():
        return None
    try:
        from docx_export import build_docx_bytes, default_filename

        data = build_docx_bytes(title, body)
        info = file_service.save_file(
            data,
            default_filename(title),
            session_id=session_id,
            description="plan_execute draft_doc",
        )
        file_id = info.get("file_id")
        if not file_id:
            return None
        preview = (body or "").strip()
        if len(preview) > 600:
            preview = preview[:600] + "…"
        return {
            "filename": info.get("original_name") or default_filename(title),
            "file_id": file_id,
            "download_url": f"/api/files/{file_id}/download",
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "title": title,
            "preview": preview,
        }
    except Exception as exc:
        print(f"[pe_tools] draft_doc docx export failed: {exc}")
        return None


def _trim(text: str, limit: int = OBS_MAX) -> str:
    text = text or ""
    if len(text) <= limit:
        return text
    return text[: limit - 20] + "\n…(truncated)"


def _citations_from_retrieve(raw: Any) -> List[Dict[str, Any]]:
    """Normalize retrieve_fn payloads (citations / law_citations / case_citations)."""
    if not isinstance(raw, dict):
        return []
    out: List[Dict[str, Any]] = []
    seen = set()
    for key in ("citations", "law_citations", "case_citations"):
        for item in raw.get(key) or []:
            if not isinstance(item, dict):
                continue
            dedupe = (
                f"{item.get('file_id') or ''}|"
                f"{item.get('document_id') or ''}|"
                f"{item.get('title') or ''}|"
                f"{item.get('article') or ''}"
            )
            if dedupe in seen:
                continue
            seen.add(dedupe)
            out.append(item)
    return out


def _observation_from_retrieve(raw: Any, scope: str) -> str:
    """Prefer formatted laws/cases text from KB retrieve_fn."""
    if not isinstance(raw, dict):
        return _trim(str(raw or ""))
    if raw.get("text"):
        return _trim(str(raw.get("text") or ""))
    if scope == "law" and raw.get("laws"):
        return _trim(str(raw.get("laws") or ""))
    if scope == "case" and raw.get("cases"):
        return _trim(str(raw.get("cases") or ""))
    parts: List[str] = []
    if raw.get("laws"):
        parts.append(str(raw.get("laws") or ""))
    if raw.get("cases"):
        parts.append(str(raw.get("cases") or ""))
    if parts:
        return _trim("\n\n".join(parts))
    if raw.get("error"):
        return _trim(f"检索失败：{raw.get('error')}")
    return _trim(str(raw))


def run_tool(name: str, args: Optional[Dict[str, Any]], ctx: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """Execute one whitelist tool. Returns {observation, citations?, artifact?}."""
    args = args or {}
    ctx = ctx or {}
    name = (name or "").strip()
    if name not in TOOL_NAMES:
        return {"observation": f"unknown tool: {name}", "citations": []}

    retrieve_fn: Optional[Callable] = ctx.get("retrieve_fn")
    write_llm = ctx.get("write_llm")
    artifact = None

    if name == "retrieve_law":
        query = str(args.get("query") or "").strip() or str(ctx.get("objective") or "")
        if not retrieve_fn:
            return {"observation": "retrieve_fn unavailable", "citations": []}
        raw = retrieve_fn(query, scopes=["law"]) or {}
        cites = _citations_from_retrieve(raw)
        out = {
            "observation": _observation_from_retrieve(raw, "law"),
            "citations": cites,
        }
        from kb_external_hint import assess_law_retrieve_miss, build_external_search_hint

        reason = assess_law_retrieve_miss(
            query,
            citations=cites,
            laws_text=str(raw.get("laws") or ""),
        )
        if reason:
            out["external_search"] = build_external_search_hint(query, reason)
        return out

    if name == "retrieve_case":
        query = str(args.get("query") or "").strip() or str(ctx.get("objective") or "")
        if not retrieve_fn:
            return {"observation": "retrieve_fn unavailable", "citations": []}
        raw = retrieve_fn(query, scopes=["case"]) or {}
        cites = _citations_from_retrieve(raw)
        # Defense in depth: drop citations that fail the same relevance gate as hits.
        try:
            from http_api_extra import prefer_hits_matching_case_query

            pseudo_hits = [
                {
                    "document": (c.get("snippet") or c.get("text") or ""),
                    "metadata": {
                        "title": c.get("title") or "",
                        "case_no": c.get("title") or "",
                        "doc_type": "case",
                        "file_id": c.get("file_id"),
                        "document_id": c.get("document_id"),
                    },
                }
                for c in cites
                if isinstance(c, dict)
            ]
            kept_meta = {
                (
                    (h.get("metadata") or {}).get("file_id") or "",
                    (h.get("metadata") or {}).get("document_id") or "",
                    (h.get("metadata") or {}).get("title") or "",
                )
                for h in prefer_hits_matching_case_query(pseudo_hits, query)
            }
            cites = [
                c
                for c in cites
                if (
                    c.get("file_id") or "",
                    c.get("document_id") or "",
                    c.get("title") or "",
                )
                in kept_meta
            ]
        except Exception:
            pass
        cases_text = str(raw.get("cases") or "")
        if cites and cases_text:
            # Rebuild observation from kept citations only (hide unrelated neighbors).
            parts = []
            for c in cites:
                title = (c.get("title") or "").strip() or "类案"
                snip = (c.get("snippet") or "").strip()
                parts.append(f"《{title}》\n{snip}" if snip else f"《{title}》")
            cases_text = "\n\n".join(parts)
        elif not cites:
            cases_text = ""
        out = {
            "observation": cases_text
            if cases_text
            else _observation_from_retrieve(
                {**raw, "cases": cases_text, "case_citations": cites}, "case"
            ),
            "citations": cites,
        }
        from kb_external_hint import (
            assess_case_retrieve_miss,
            build_case_external_search_hint,
        )

        reason = assess_case_retrieve_miss(
            query, citations=cites, cases_text=cases_text
        )
        if reason:
            out["observation"] = (
                "本地知识库未命中相关类案（已过滤与查询关键词不相关的结果）。"
                "请换更贴切的检索词，或补充类案入库后再试。"
            )
            out["citations"] = []
            out["external_search"] = build_case_external_search_hint(query, reason)
        return out

    if name == "read_evidence":
        from case_materials import get_case_evidence_text, resolve_cases_for_file_id

        file_id = args.get("file_id")
        scope = str(
            ctx.get("case_scope")
            or ("single" if ctx.get("case_id") not in (None, "") else "none")
        )
        case_store = ctx.get("case_store")
        file_service = ctx.get("file_service")
        case_id = ctx.get("case_id")

        if scope == "none" or (case_id in (None, "", "*") and scope != "all_permitted"):
            if scope != "all_permitted":
                return {
                    "observation": (
                        "当前未选择案件，无法读取证据材料。"
                        "请在下拉框选择具体案件或「全选（我有权限的案件）」，"
                        "或直接补充原告/被告等当事人信息后继续；"
                        "若坚持不绑定案件，请说明「先按占位起草」。"
                    ),
                    "citations": [],
                }

        if scope == "all_permitted":
            arg_case = args.get("case_id")
            permitted = [int(x) for x in (ctx.get("permitted_case_ids") or [])]
            if arg_case not in (None, ""):
                try:
                    cid = int(arg_case)
                except (TypeError, ValueError):
                    return {"observation": "read_evidence 的 case_id 无效", "citations": []}
                if cid not in permitted:
                    return {"observation": "无权在该案件下读取证据", "citations": []}
                case_id = cid
            else:
                if not file_id:
                    return {
                        "observation": "全选模式下请提供 file_id，或指定案件 case_id",
                        "citations": [],
                    }
                hits = resolve_cases_for_file_id(str(file_id), permitted, case_store)
                if len(hits) == 0:
                    return {
                        "observation": "在有权限的案件中未找到该证据文件，请确认 file_id 或改选具体案件",
                        "citations": [],
                    }
                if len(hits) > 1:
                    return {
                        "observation": f"该文件关联多个案件 {hits}，请指定 case_id 后再读取",
                        "citations": [],
                    }
                case_id = hits[0]

        if not (file_id and case_id not in (None, "") and case_store and file_service):
            return {
                "observation": (
                    "读取证据缺少必要参数。请选择案件并提供 file_id，或改用「全选」后提供可定位的文件标识。"
                ),
                "citations": [],
            }
        try:
            text = get_case_evidence_text(int(case_id), str(file_id), case_store, file_service)
        except Exception as exc:
            return {"observation": f"读取证据失败：{exc}", "citations": []}
        return {"observation": _trim(text or "(empty)"), "citations": []}

    if name == "draft_doc":
        prompt = str(args.get("prompt") or args.get("instruction") or ctx.get("objective") or "")
        case_context = str(ctx.get("case_context") or "")
        if case_context:
            from case_materials import format_user_with_case_context

            prompt = format_user_with_case_context(prompt, case_context)
        query = str(ctx.get("objective") or prompt)
        kb_store = ctx.get("kb_store")
        file_service = ctx.get("file_service")
        session_id = ctx.get("session_id")
        hit = None
        if kb_store is not None:
            from kb_template_resolve import match_template

            hit = match_template(kb_store, query)

        if hit and file_service:
            raw = _load_template_docx_bytes(file_service, hit.get("file_id"))
            tmpl_name = hit.get("name") or "模版"
            if not raw:
                return {
                    "observation": f"已匹配模版《{tmpl_name}》，但无法读取模版文件，请检查知识库文件后重试。",
                    "citations": [],
                }
            from docx_form_fill import (
                build_element_dict,
                fill_docx_bytes,
                has_fillable_slots,
                scan_slots_from_document,
            )

            if has_fillable_slots(raw):
                import io

                from docx import Document

                doc = Document(io.BytesIO(raw))
                slots = scan_slots_from_document(doc)
                slot_keys = [s["key"] for s in slots]
                elements = build_element_dict(
                    prompt,
                    slot_keys=slot_keys,
                    write_llm=write_llm,
                )
                filled = fill_docx_bytes(raw, elements)
                filename = f"{tmpl_name}.docx"
                try:
                    info = file_service.save_file(
                        filled,
                        filename,
                        session_id=session_id,
                        description="plan_execute draft_doc",
                    )
                except Exception as exc:
                    print(f"[pe_tools] draft_doc template fill save failed: {exc}")
                    return {
                        "observation": f"模版《{tmpl_name}》填表完成但保存失败，请重试。",
                        "citations": [],
                    }
                if not (isinstance(info, dict) and info.get("file_id")):
                    return {
                        "observation": f"模版《{tmpl_name}》填表完成但保存失败，请重试。",
                        "citations": [],
                    }
                filled_summary = "、".join(
                    f"{k}={v}"
                    for k, v in elements.items()
                    if v and str(v).strip()
                )
                obs = (
                    f"已根据模版《{tmpl_name}》填写要素式文书，请下载核阅。"
                    f"缺项已标【待补充】。"
                )
                if filled_summary:
                    obs += f"\n已填要素：{filled_summary}"
                preview = _preview_text_from_docx_bytes(filled) or obs
                artifact = _artifact_from_saved_docx(
                    info,
                    title=tmpl_name,
                    preview=preview,
                )
                return {"observation": _trim(obs), "citations": [], "artifact": artifact}

            if not write_llm:
                return {"observation": "write_llm unavailable for draft_doc", "citations": []}
            from kb_template_resolve import resolve_template_text

            template_text, _, _ = resolve_template_text(
                tmpl_name,
                kb_store=kb_store,
                file_service=file_service,
                vector_service=ctx.get("vector_service"),
            )
            if not template_text:
                get_text = getattr(file_service, "get_file_text", None)
                if callable(get_text):
                    template_text = (get_text(hit.get("file_id")) or "").strip()
            skills = ctx.get("skills")
            system = _draft_doc_system_prompt(skills if isinstance(skills, list) else [])
            system += f"\n已匹配知识库模版《{tmpl_name}》（非表格填槽），请参照模版结构与表述起草完整文书正文。"
            if template_text:
                system += f"\n【模版参考，禁止原样照抄未改写段落】\n{template_text[:8000]}"
            try:
                body = write_llm(system, prompt, ctx.get("messages") or []) or ""
            except Exception as exc:
                return {"observation": f"draft_doc failed: {exc}", "citations": []}
            title = str(args.get("title") or ctx.get("doc_title") or "").strip()
            if not title:
                title = _infer_draft_title(str(prompt or "") + " " + str(ctx.get("objective") or ""))
            artifact = _export_docx_artifact(title, body, file_service, session_id=session_id)
            obs = f"已套用模版《{tmpl_name}》（非表格填槽）。\n\n{body}"
            return {"observation": _trim(obs), "citations": [], "artifact": artifact}

        if not write_llm:
            return {"observation": "write_llm unavailable for draft_doc", "citations": []}
        skills = ctx.get("skills")
        if not skills:
            objective = str(ctx.get("objective") or prompt or "").strip()
            match_fn = ctx.get("skill_match_fn")
            if callable(match_fn) and objective:
                try:
                    skills = match_fn(objective) or []
                except Exception:
                    skills = []
            elif objective:
                try:
                    from skill_service import SkillService
                    import os

                    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
                    # server/agents -> project root
                    project = os.path.dirname(root)
                    skills = SkillService(os.path.join(project, "skills")).match(objective, limit=3)
                except Exception:
                    skills = []
        system = _draft_doc_system_prompt(skills if isinstance(skills, list) else [])
        try:
            body = write_llm(system, prompt, ctx.get("messages") or []) or ""
        except Exception as exc:
            return {"observation": f"draft_doc failed: {exc}", "citations": []}
        title = str(args.get("title") or ctx.get("doc_title") or "").strip()
        if not title:
            title = _infer_draft_title(
                str(prompt or "") + " " + str(ctx.get("objective") or "")
            )
        artifact = _export_docx_artifact(
            title,
            body,
            file_service,
            session_id=session_id,
        )
        observation = body
        # Only prefix when kb_store was present but template match failed (backward compat).
        if kb_store is not None and hit is None:
            observation = f"未命中模版库，已按说明自由起草。\n\n{observation}"
        return {"observation": _trim(observation), "citations": [], "artifact": artifact}

    if name == "reason":
        prompt = str(args.get("prompt") or args.get("instruction") or ctx.get("objective") or "")
        case_context = str(ctx.get("case_context") or "")
        if case_context:
            from case_materials import format_user_with_case_context

            prompt = format_user_with_case_context(prompt, case_context)
        if not write_llm:
            return {"observation": "(no write_llm) " + prompt, "citations": []}
        system = "你是法律分析助手。根据给定步骤要求给出简明推理或结论，引用须基于已提供材料。"
        try:
            body = write_llm(system, prompt, ctx.get("messages") or []) or ""
        except Exception as exc:
            return {"observation": f"reason failed: {exc}", "citations": []}
        return {"observation": _trim(body), "citations": []}

    return {"observation": f"unhandled tool: {name}", "citations": []}
