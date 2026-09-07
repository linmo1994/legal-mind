"""Scan and fill 要素式 Word tables (placeholders + label→right cell)."""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional

PLACEHOLDER_MISSING = "【待补充】"

# Normalized label text → canonical element key
LABEL_KEYS = {
    "原告": "原告",
    "被告": "被告",
    "申请人": "申请人",
    "被申请人": "被申请人",
    "法定代表人": "法定代表人",
    "委托诉讼代理人": "委托诉讼代理人",
    "身份证号": "身份证号",
    "身份证号码": "身份证号",
    "住所地": "住所地",
    "住址": "住所地",
    "联系电话": "联系电话",
    "诉讼请求": "诉讼请求",
    "事实与理由": "事实与理由",
    "事实和理由": "事实与理由",
}

# Longest-first: {{...}} before {...} so double-brace placeholders are not
# partially matched as single-brace.
_PLACEHOLDER_RES = (
    re.compile(r"\{\{([^{}]+)\}\}"),
    re.compile(r"\{([^{}]+)\}"),
    re.compile(r"【([^【】]+)】"),
)


def _norm_label(text: str) -> str:
    t = re.sub(r"\s+", "", (text or "").strip())
    t = t.strip("：:．.、")
    return t


def _cell_empty(text: str) -> bool:
    t = (text or "").strip()
    if not t:
        return True
    if re.fullmatch(r"[_＿…\-\s]+", t):
        return True
    if t in ("【待补充】", "（待填写）", "(待填写)"):
        return True
    return False


def scan_slots_from_document(doc) -> List[Dict[str, Any]]:
    slots: List[Dict[str, Any]] = []
    seen = set()
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            cells = row.cells
            for ci, cell in enumerate(cells):
                text = cell.text or ""
                for rx in _PLACEHOLDER_RES:
                    for m in rx.finditer(text):
                        key = _norm_label(m.group(1))
                        if not key or key in ("待补充",):
                            continue
                        # Skip if key is only a section header style placeholder duplicate
                        sig = ("placeholder", ti, ri, ci, key)
                        if sig in seen:
                            continue
                        seen.add(sig)
                        slots.append(
                            {
                                "key": key,
                                "mode": "placeholder",
                                "table_i": ti,
                                "row_i": ri,
                                "cell_i": ci,
                                "raw": m.group(0),
                            }
                        )
            # label → right cell (2+ cols)
            if len(cells) >= 2:
                left = _norm_label(cells[0].text or "")
                key = LABEL_KEYS.get(left)
                if key and _cell_empty(cells[1].text or ""):
                    sig = ("label_right", ti, ri, 1, key)
                    if sig not in seen:
                        seen.add(sig)
                        slots.append(
                            {
                                "key": key,
                                "mode": "label_right",
                                "table_i": ti,
                                "row_i": ri,
                                "cell_i": 1,
                                "label": left,
                            }
                        )
    return slots


def _set_cell_text(cell, value: str) -> None:
    value = value if value is not None else ""
    # Clear paragraphs then set first paragraph text (simple, preserves table)
    if cell.paragraphs:
        cell.paragraphs[0].text = value
        for p in cell.paragraphs[1:]:
            p.text = ""
    else:
        cell.text = value


def fill_document(doc, elements: Dict[str, Any], slots: Optional[List[Dict[str, Any]]] = None):
    slots = slots if slots is not None else scan_slots_from_document(doc)
    elements = elements or {}
    for slot in slots:
        key = slot["key"]
        raw_val = elements.get(key)
        if raw_val is None or str(raw_val).strip() == "":
            val = PLACEHOLDER_MISSING
        else:
            val = str(raw_val).strip()
        ti, ri, ci = slot["table_i"], slot["row_i"], slot["cell_i"]
        if ti >= len(doc.tables):
            continue
        table = doc.tables[ti]
        if ri >= len(table.rows):
            continue
        row = table.rows[ri]
        if ci >= len(row.cells):
            continue
        cell = row.cells[ci]
        if slot["mode"] == "placeholder":
            text = cell.text or ""
            raw = slot.get("raw") or ""
            if raw and raw in text:
                _set_cell_text(cell, text.replace(raw, val, 1))
            else:
                # replace any {key} / 【key】 occurrence
                new_text = text
                for rx in _PLACEHOLDER_RES:
                    new_text = rx.sub(
                        lambda m: val if _norm_label(m.group(1)) == key else m.group(0),
                        new_text,
                        count=1,
                    )
                _set_cell_text(cell, new_text)
        else:
            _set_cell_text(cell, val)
    return doc


def fill_docx_bytes(docx_bytes: bytes, elements: Dict[str, Any]) -> bytes:
    from docx import Document

    doc = Document(io.BytesIO(docx_bytes or b""))
    fill_document(doc, elements)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def has_fillable_slots(docx_bytes: bytes) -> bool:
    from docx import Document

    try:
        doc = Document(io.BytesIO(docx_bytes or b""))
    except Exception:
        return False
    return bool(scan_slots_from_document(doc))
