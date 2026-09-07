import io
import unittest
from docx import Document

from docx_form_fill import (
    PLACEHOLDER_MISSING,
    fill_docx_bytes,
    scan_slots_from_document,
)


def _doc_with_placeholder_and_label():
    doc = Document()
    t1 = doc.add_table(rows=1, cols=1)
    t1.cell(0, 0).text = "原告：{原告姓名}"
    t2 = doc.add_table(rows=1, cols=2)
    t2.cell(0, 0).text = "被告"
    t2.cell(0, 1).text = ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDocxFormFill(unittest.TestCase):
    def test_scan_finds_placeholder_and_label_right(self):
        raw = _doc_with_placeholder_and_label()
        doc = Document(io.BytesIO(raw))
        slots = scan_slots_from_document(doc)
        keys = {s["key"] for s in slots}
        self.assertIn("原告姓名", keys)
        self.assertIn("被告", keys)
        modes = {s["key"]: s["mode"] for s in slots}
        self.assertEqual(modes["原告姓名"], "placeholder")
        self.assertEqual(modes["被告"], "label_right")

    def test_fill_writes_values_and_missing_marker(self):
        raw = _doc_with_placeholder_and_label()
        out = fill_docx_bytes(raw, {"原告姓名": "张三"})
        doc = Document(io.BytesIO(out))
        self.assertIn("张三", doc.tables[0].cell(0, 0).text)
        self.assertIn(PLACEHOLDER_MISSING, doc.tables[1].cell(0, 1).text)
        self.assertNotIn("{原告姓名}", doc.tables[0].cell(0, 0).text)
