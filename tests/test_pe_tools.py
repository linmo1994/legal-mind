#!/usr/bin/env python3
import io
import os
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "server"))

from agents.pe_tools import TOOL_NAMES, run_tool  # noqa: E402


class TestPeTools(unittest.TestCase):
    def test_tool_names(self):
        self.assertEqual(
            set(TOOL_NAMES),
            {"retrieve_law", "retrieve_case", "read_evidence", "draft_doc", "reason"},
        )

    def test_retrieve_law_calls_retrieve_fn(self):
        seen = []

        def retrieve(query, scopes=None):
            seen.append((query, tuple(scopes or ())))
            return {"text": "法条摘要", "citations": [{"title": "民法典", "article": "第667条"}]}

        out = run_tool(
            "retrieve_law",
            {"query": "民间借贷利率"},
            ctx={"retrieve_fn": retrieve, "write_llm": None},
        )
        self.assertEqual(seen, [("民间借贷利率", ("law",))])
        self.assertIn("法条摘要", out["observation"])
        self.assertEqual(len(out.get("citations") or []), 1)

    def test_retrieve_law_reads_kb_law_citations(self):
        """KB retrieve_fn returns laws + law_citations (not citations/text)."""

        def retrieve(query, scopes=None):
            return {
                "laws": "民法典第六百六十七条……",
                "cases": "",
                "law_citations": [
                    {
                        "title": "民法典",
                        "article": "第六百六十七条",
                        "file_id": "f1",
                        "doc_type": "law",
                    }
                ],
                "case_citations": [],
            }

        out = run_tool(
            "retrieve_law",
            {"query": "借款合同"},
            ctx={"retrieve_fn": retrieve, "write_llm": None},
        )
        self.assertIn("民法典", out["observation"])
        self.assertEqual(len(out["citations"]), 1)
        self.assertEqual(out["citations"][0]["file_id"], "f1")
        self.assertNotIn("{", out["observation"][:1])  # not raw dict dump

    def test_unknown_tool(self):
        out = run_tool("nope", {}, ctx={})
        self.assertIn("unknown", out["observation"].lower())

    def test_reason_uses_write_llm(self):
        def write_llm(system, user, hist=None):
            return "推理结论"

        out = run_tool(
            "reason",
            {"prompt": "分析利率是否合法"},
            ctx={"write_llm": write_llm},
        )
        self.assertEqual(out["observation"], "推理结论")

    @patch("case_materials.get_case_evidence_text")
    def test_read_evidence_passes_str_file_id(self, mock_get_text):
        mock_get_text.return_value = "证据正文"
        out = run_tool(
            "read_evidence",
            {"file_id": 42},
            ctx={
                "case_id": 1,
                "case_store": object(),
                "file_service": object(),
            },
        )
        mock_get_text.assert_called_once()
        self.assertIsInstance(mock_get_text.call_args[0][1], str)
        self.assertEqual(mock_get_text.call_args[0][1], "42")
        self.assertIn("证据正文", out["observation"])

    def test_read_evidence_no_case_chinese_hint(self):
        out = run_tool(
            "read_evidence",
            {"file_id": "f1"},
            ctx={"case_scope": "none", "case_id": None},
        )
        self.assertIn("选择案件", out["observation"])
        self.assertIn("全选", out["observation"])
        self.assertNotIn("case_store", out["observation"])

    @patch("case_materials.get_case_evidence_text")
    def test_read_evidence_all_scope_resolves_unique(self, mock_get):
        mock_get.return_value = "全文A"
        store = type("S", (), {})()

        def get_case(cid):
            if cid == 7:
                return {"id": 7, "meta": {"evidence_file_ids": ["f1"]}}
            return {"id": cid, "meta": {"evidence_file_ids": []}}

        store.get_case = get_case
        out = run_tool(
            "read_evidence",
            {"file_id": "f1"},
            ctx={
                "case_scope": "all_permitted",
                "permitted_case_ids": [7, 8],
                "case_store": store,
                "file_service": object(),
            },
        )
        mock_get.assert_called_once()
        self.assertEqual(mock_get.call_args[0][0], 7)
        self.assertIn("全文A", out["observation"])

    def test_read_evidence_all_scope_ambiguous(self):
        store = type("S", (), {})()
        store.get_case = lambda cid: {
            "id": cid,
            "meta": {"evidence_file_ids": ["f1"]},
        }
        out = run_tool(
            "read_evidence",
            {"file_id": "f1"},
            ctx={
                "case_scope": "all_permitted",
                "permitted_case_ids": [1, 2],
                "case_store": store,
                "file_service": object(),
            },
        )
        self.assertIn("多个案件", out["observation"])

    def test_retrieve_law_miss_adds_external_search(self):
        def retrieve(q, scopes=None):
            return {"laws": "", "law_citations": []}

        out = run_tool(
            "retrieve_law",
            {"query": "某某不存在法第一条"},
            {"retrieve_fn": retrieve, "objective": "x"},
        )
        self.assertTrue(out.get("external_search", {}).get("needed"))
        self.assertEqual(out["external_search"]["provider"], "npc_flk")

    def test_retrieve_law_hit_no_external_search(self):
        def retrieve(q, scopes=None):
            return {
                "laws": "《劳动合同法》\n第六十四条……",
                "law_citations": [
                    {
                        "title": "中华人民共和国劳动合同法",
                        "article": "第六十四条",
                        "snippet": "第六十四条　被派遣劳动者……",
                    }
                ],
            }

        out = run_tool(
            "retrieve_law",
            {"query": "劳动合同法第六十四条"},
            {"retrieve_fn": retrieve},
        )
        self.assertFalse(out.get("external_search"))

    def test_retrieve_case_filters_irrelevant_citations(self):
        def retrieve(q, scopes=None):
            return {
                "cases": "两案…",
                "case_citations": [
                    {
                        "title": "（2023）粤2071民初27790号",
                        "file_id": "f-food",
                        "document_id": "kb_food",
                        "snippet": "餐饮服务合同纠纷，主张十倍赔偿",
                    },
                    {
                        "title": "（2025）最高法民再142号",
                        "file_id": "f-loan",
                        "document_id": "kb_loan",
                        "snippet": "民间借贷担保纠纷保证责任",
                    },
                ],
            }

        out = run_tool(
            "retrieve_case",
            {"query": "查找餐饮服务合同违约10倍赔偿的案例"},
            {"retrieve_fn": retrieve},
        )
        cites = out.get("citations") or []
        self.assertEqual(len(cites), 1)
        self.assertIn("粤2071", cites[0]["title"])
        self.assertNotIn("最高法民再142", out.get("observation") or "")
        self.assertFalse(out.get("external_search"))

    def test_retrieve_case_miss_adds_external_search(self):
        def retrieve(q, scopes=None):
            return {"cases": "", "case_citations": []}

        out = run_tool(
            "retrieve_case",
            {"query": "食品服务合同 十倍赔偿"},
            {"retrieve_fn": retrieve},
        )
        self.assertEqual(out.get("citations") or [], [])
        self.assertIn("未命中相关类案", out["observation"])
        self.assertTrue(out.get("external_search", {}).get("needed"))
        self.assertEqual(out["external_search"]["provider"], "court_wenshu")

    def test_retrieve_case_hit_no_external_search(self):
        def retrieve(q, scopes=None):
            return {
                "cases": "（2025）最高法民再142号\n民间借贷……",
                "case_citations": [
                    {
                        "title": "（2025）最高法民再142号",
                        "file_id": "f1",
                        "snippet": "民间借贷担保纠纷……",
                    }
                ],
            }

        out = run_tool(
            "retrieve_case",
            {"query": "民间借贷担保"},
            {"retrieve_fn": retrieve},
        )
        self.assertFalse(out.get("external_search"))
        self.assertEqual(len(out.get("citations") or []), 1)

    def test_draft_doc_injects_doc_writing_skills(self):
        seen = {}

        def write_llm(system, user, hist=None):
            seen["system"] = system
            seen["user"] = user
            return "民事起诉状\n……\n此致\n某某人民法院\n具状人：张三\n日期：____"

        skills = [
            {
                "id": "legal-pleading",
                "name": "法律文书起草",
                "applies_to": ["doc_writing", "orchestrator"],
                "body": "结尾写「此致 / 人民法院」及具状人、日期占位。只输出文书正文。",
            },
            {
                "id": "other",
                "name": "无关技能",
                "applies_to": ["text_analysis"],
                "body": "不应出现在文书系统提示",
            },
        ]
        out = run_tool(
            "draft_doc",
            {"prompt": "生成民间借贷起诉状"},
            {"write_llm": write_llm, "skills": skills, "objective": "生成起诉状"},
        )
        self.assertIn("起诉状", out["observation"] or "")
        sys = seen.get("system") or ""
        self.assertIn("法律文书起草", sys)
        self.assertIn("此致", sys)
        self.assertIn("【内部技能", sys)
        self.assertNotIn("无关技能", sys)
        self.assertIn("文书式", sys)

    def test_draft_doc_without_skills_still_requires_formal_closing(self):
        seen = {}

        def write_llm(system, user, hist=None):
            seen["system"] = system
            return "草稿"

        run_tool("draft_doc", {"prompt": "写协议"}, {"write_llm": write_llm})
        self.assertIn("落款", seen.get("system") or "")


    def test_draft_doc_exports_docx_artifact(self):
        saved = {}

        def write_llm(system, user, hist=None):
            return "民事起诉状\n原告：张三\n此致\n人民法院\n具状人：张三"

        class FakeFS:
            def save_file(self, data, filename, session_id=None, description=None):
                saved["data"] = data
                saved["filename"] = filename
                saved["session_id"] = session_id
                return {"file_id": "fid-docx-1", "original_name": filename}

        out = run_tool(
            "draft_doc",
            {"prompt": "生成民间借贷起诉状"},
            {
                "write_llm": write_llm,
                "file_service": FakeFS(),
                "session_id": "sess-1",
                "objective": "生成起诉状",
            },
        )
        art = out.get("artifact") or {}
        self.assertEqual(art.get("file_id"), "fid-docx-1")
        self.assertTrue(str(art.get("filename") or "").endswith(".docx"))
        self.assertIn("download", art.get("download_url") or "")
        self.assertTrue(isinstance(saved.get("data"), (bytes, bytearray)))
        self.assertGreater(len(saved["data"]), 20)
        self.assertEqual(saved.get("session_id"), "sess-1")

    def test_draft_doc_fills_matched_element_template(self):
        doc = Document()
        t1 = doc.add_table(rows=1, cols=1)
        t1.cell(0, 0).text = "原告：{原告}"
        t2 = doc.add_table(rows=1, cols=2)
        t2.cell(0, 0).text = "被告"
        t2.cell(0, 1).text = ""
        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        doc.save(tmp.name)
        tmp.close()
        template_path = tmp.name
        self.addCleanup(lambda: os.path.exists(template_path) and os.unlink(template_path))

        saved = {}

        class FakeKb:
            def list_documents(self, doc_type="template", limit=200, offset=0):
                return [
                    {
                        "id": "t1",
                        "doc_type": "template",
                        "file_id": "f-template",
                        "title": "民间借贷纠纷起诉状",
                        "status": "ready",
                        "meta": {
                            "template_name": "民间借贷纠纷起诉状",
                            "document_type": "起诉状",
                            "validity": "有效",
                        },
                    }
                ]

        class FakeFS:
            def get_file(self, file_id):
                return {"file_path": template_path, "file_id": file_id}

            def save_file(self, data, filename, session_id=None, description=None):
                saved["data"] = data
                saved["filename"] = filename
                saved["session_id"] = session_id
                return {"file_id": "filled-docx-1", "original_name": filename}

        def write_llm(system, user, hist=None):
            return ""

        out = run_tool(
            "draft_doc",
            {"prompt": "请写民间借贷纠纷起诉状，原告张三被告李四"},
            {
                "write_llm": write_llm,
                "kb_store": FakeKb(),
                "file_service": FakeFS(),
                "objective": "请写民间借贷纠纷起诉状",
                "session_id": "sess-tmpl",
            },
        )
        self.assertIn("模版", out["observation"])
        self.assertIn("要素式", out["observation"])
        self.assertIn("请下载核阅", out["observation"])
        self.assertIn("【待补充】", out["observation"])
        art = out.get("artifact") or {}
        self.assertEqual(art.get("file_id"), "filled-docx-1")
        self.assertTrue(isinstance(saved.get("data"), (bytes, bytearray)))
        filled = Document(io.BytesIO(saved["data"]))
        self.assertIn("张三", filled.tables[0].cell(0, 0).text)
        self.assertIn("李四", filled.tables[1].cell(0, 1).text)

    def test_draft_doc_save_failure_after_fill(self):
        doc = Document()
        t = doc.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "原告：{原告}"
        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        doc.save(tmp.name)
        tmp.close()
        template_path = tmp.name
        self.addCleanup(lambda: os.path.exists(template_path) and os.unlink(template_path))

        class FakeKb:
            def list_documents(self, doc_type="template", limit=200, offset=0):
                return [
                    {
                        "id": "t1",
                        "doc_type": "template",
                        "file_id": "f-template",
                        "title": "民间借贷纠纷起诉状",
                        "status": "ready",
                        "meta": {
                            "template_name": "民间借贷纠纷起诉状",
                            "document_type": "起诉状",
                            "validity": "有效",
                        },
                    }
                ]

        class FakeFS:
            def get_file(self, file_id):
                return {"file_path": template_path, "file_id": file_id}

            def save_file(self, data, filename, session_id=None, description=None):
                raise RuntimeError("disk full")

        out = run_tool(
            "draft_doc",
            {"prompt": "请写民间借贷纠纷起诉状，原告张三"},
            {
                "write_llm": lambda s, u, h=None: "不应走到自由起草",
                "kb_store": FakeKb(),
                "file_service": FakeFS(),
                "objective": "请写民间借贷纠纷起诉状",
            },
        )
        self.assertIn("保存失败", out["observation"])
        self.assertNotIn("不应走到自由起草", out["observation"])
        self.assertIsNone(out.get("artifact"))

    def test_draft_doc_unmatched_mentions_fallback(self):
        class FakeKb:
            def list_documents(self, doc_type="template", limit=200, offset=0):
                return [
                    {
                        "id": "t2",
                        "doc_type": "template",
                        "file_id": "f2",
                        "title": "劳动争议仲裁申请书",
                        "status": "ready",
                        "meta": {
                            "template_name": "劳动争议仲裁申请书",
                            "document_type": "申请书",
                            "validity": "有效",
                        },
                    }
                ]

        class FakeFS:
            def save_file(self, data, filename, session_id=None, description=None):
                return {"file_id": "fid-free", "original_name": filename}

        def write_llm(system, user, hist=None):
            return "自由起草正文"

        out = run_tool(
            "draft_doc",
            {"prompt": "随便写点闲聊"},
            {
                "write_llm": write_llm,
                "kb_store": FakeKb(),
                "file_service": FakeFS(),
            },
        )
        self.assertIn("未命中模版", out["observation"])
        self.assertIn("自由起草正文", out["observation"])


if __name__ == "__main__":
    unittest.main()